# -*- coding: utf-8 -*-
"""
RISCO DE FALTA - servico proprio (separado do painel Shopee/ML).

Cruza as saidas do Bling com a planilha de estoque do fornecedor e aponta o que
acaba em 1 semana / 1 mes.  Tres abas:
  - Saidas por NF          -> NFs de saida (nao usa pedido de venda)
  - Full (pedidos de venda)-> pedidos de venda das lojas do Full (no Bling nao geram NF por venda)
  - Total                  -> NF + Full somados

Variaveis de ambiente (Render):
  BLING_CLIENT_ID, BLING_CLIENT_SECRET   app privado do Bling criado SO para este servico
  APP_SENHA                              senha de entrada da pagina
  SECRET_KEY                             texto aleatorio longo (sessao)
  DATABASE_URL                           (recomendado) string de conexao do Neon; guarda token e cache
  FULL_LOJA_IDS                          (opcional) ids das lojas do Full separados por virgula
Start command no Render:  gunicorn app:app --workers 1 --threads 4 --timeout 120
"""
import io
import json
import math
import os
import re
import secrets
import hmac
import threading
import time
import datetime as dt
from collections import Counter, defaultdict
from functools import wraps

import requests
from flask import Flask, request, jsonify, session, redirect, send_file, Response

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(24)

CID = os.environ.get("BLING_CLIENT_ID", "")
CSEC = os.environ.get("BLING_CLIENT_SECRET", "")
SENHA = os.environ.get("APP_SENHA", "")
DB_URL = os.environ.get("DATABASE_URL", "")
API = "https://api.bling.com.br/Api/v3"
OAUTH = "https://www.bling.com.br/Api/v3/oauth"
SIT_NF_OK = {5, 6, 7}            # Autorizada, DANFE emitida, Registrada
SIT_PEDIDO_CANCELADO = 2         # situacao.valor = 2 -> cancelado
PAUSA = 0.40                     # Bling aceita ~3 requisicoes por segundo
FULL_PADRAO = os.environ.get("FULL_LOJA_IDS", "206006900")


# =============================================================================
# Armazenamento simples chave/valor (Neon se houver DATABASE_URL, senao arquivo)
# =============================================================================
_kv_lock = threading.Lock()
ARQ_KV = os.environ.get("KV_ARQUIVO", "/tmp/risco_falta_kv.json")


def _sql(sql, args=()):
    import psycopg2
    c = psycopg2.connect(DB_URL, connect_timeout=10)
    try:
        with c, c.cursor() as cur:
            cur.execute(sql, args)
            return cur.fetchone() if cur.description else None
    finally:
        c.close()


def kv_get(k, padrao=None):
    try:
        if DB_URL:
            _sql("create table if not exists kv (k text primary key, v text)")
            r = _sql("select v from kv where k=%s", (k,))
            return json.loads(r[0]) if r else padrao
        with _kv_lock:
            d = json.load(open(ARQ_KV)) if os.path.exists(ARQ_KV) else {}
        return d.get(k, padrao)
    except Exception:
        return padrao


def kv_set(k, v):
    if DB_URL:
        _sql("create table if not exists kv (k text primary key, v text)")
        _sql("insert into kv(k, v) values (%s, %s) on conflict (k) do update set v = excluded.v",
             (k, json.dumps(v)))
        return
    with _kv_lock:
        d = json.load(open(ARQ_KV)) if os.path.exists(ARQ_KV) else {}
        d[k] = v
        json.dump(d, open(ARQ_KV, "w"))


# =============================================================================
# Bling: OAuth + chamadas
# =============================================================================
_tok_lock = threading.Lock()
_TOK = {}


def _pedir_token(data):
    r = requests.post(OAUTH + "/token", data=data, auth=(CID, CSEC), timeout=30)
    if r.status_code >= 400:
        raise RuntimeError("Bling recusou o token: %s" % r.text[:200])
    d = r.json()
    d["expira_em"] = time.time() + int(d.get("expires_in", 21600)) - 120
    kv_set("bling_token", d)
    _TOK.clear()
    _TOK.update(d)
    return d


def _access_token(forcar=False):
    with _tok_lock:
        if not _TOK:
            _TOK.update(kv_get("bling_token") or {})
        if not _TOK.get("refresh_token"):
            raise RuntimeError("Bling não conectado: clique em Conectar ao Bling.")
        if forcar or time.time() >= _TOK.get("expira_em", 0):
            _pedir_token({"grant_type": "refresh_token", "refresh_token": _TOK["refresh_token"]})
        return _TOK["access_token"]


def bling_get(path, params=None):
    tk = _access_token()
    for tentativa in range(2):
        r = requests.get(API + path, params=params or {}, headers={"Authorization": "Bearer " + tk}, timeout=40)
        if r.status_code == 401 and tentativa == 0:
            tk = _access_token(forcar=True)
            continue
        if r.status_code == 429:
            raise RuntimeError("429 limite de requisições do Bling")
        r.raise_for_status()
        return r.json()


def _chamar(path, params=None, tentativas=4):
    ultimo = None
    for i in range(tentativas):
        try:
            return bling_get(path, params)
        except Exception as e:
            ultimo = e
            time.sleep(1.5 * (i + 1))
    raise ultimo


# =============================================================================
# Login
# =============================================================================
def exige_login(f):
    @wraps(f)
    def w(*a, **k):
        if not SENHA:
            return Response("Defina a variável APP_SENHA no Render.", 500)
        if not session.get("ok"):
            if request.path.startswith("/api/"):
                return jsonify(ok=False, erro="Sessão expirada. Entre novamente."), 401
            return redirect("/login")
        return f(*a, **k)
    return w


LOGIN_HTML = """<!doctype html><html lang="pt-br"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Risco de falta</title><body style="font:16px system-ui;background:#f5f6f8;display:grid;place-items:center;height:100vh;margin:0">
<form method="post" style="background:#fff;border:1px solid #dfe3ea;border-radius:8px;padding:24px;width:280px">
<h2 style="margin:0 0 14px">Risco de falta</h2>
<input type="password" name="senha" placeholder="Senha" autofocus style="width:100%;padding:9px;box-sizing:border-box;border:1px solid #c9cfda;border-radius:6px">
<button style="margin-top:12px;width:100%;padding:10px;border:0;border-radius:6px;background:#1f5eff;color:#fff;font-size:15px">Entrar</button>
<div style="color:#b3261e;margin-top:10px;font-size:14px">{{ERRO}}</div></form></body></html>"""


@app.route("/login", methods=["GET", "POST"])
def login():
    erro = ""
    if request.method == "POST":
        if SENHA and hmac.compare_digest(request.form.get("senha", ""), SENHA):
            session["ok"] = True
            return redirect("/")
        erro = "Senha incorreta."
    return Response(LOGIN_HTML.replace("{{ERRO}}", erro), mimetype="text/html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/login")


@app.route("/bling/conectar")
@exige_login
def bling_conectar():
    st = secrets.token_urlsafe(16)
    session["oauth_state"] = st
    return redirect("%s/authorize?response_type=code&client_id=%s&state=%s" % (OAUTH, CID, st))


@app.route("/bling/callback")
@exige_login
def bling_callback():
    if request.args.get("state") != session.get("oauth_state"):
        return Response("State inválido. Volte e clique em Conectar ao Bling de novo.", 400)
    try:
        _pedir_token({"grant_type": "authorization_code", "code": request.args.get("code", "")})
    except Exception as e:
        return Response("Falha ao conectar: %s" % e, 400)
    return redirect("/")


# =============================================================================
# Dados em memoria (e salvos no kv)
# =============================================================================
CACHE = {"nf": kv_get("cache_nf", {}) or {}, "full": kv_get("cache_full", {}) or {}}
PLAN = kv_get("planilha", {}) or {}
PLAN_NOME = kv_get("planilha_nome", "")
LOJAS_VISTAS = {}
ULTIMA = {"linhas": [], "aba": ""}
_lock = threading.Lock()
JOB = {"fonte": "", "estado": "parado", "feitas": 0, "total": 0, "erro": ""}


def _agora():
    return dt.datetime.utcnow() - dt.timedelta(hours=3)


def _norm_codigo(cod):
    return re.sub(r"\(\s*(\w+)\s*\)", r"(\1)", str(cod).strip().upper())


def _itens(resp):
    out = []
    for it in ((resp or {}).get("data", {}) or {}).get("itens", []) or []:
        cod = _norm_codigo(it.get("codigo", ""))
        try:
            q = float(it.get("quantidade", 0))
        except (TypeError, ValueError):
            q = 0
        if cod and q:
            out.append([cod, q])
    return out


def _salvar_cache(fonte):
    try:
        kv_set("cache_" + fonte, CACHE[fonte])
    except Exception:
        pass


# -------------------------- sincronizacao: NF de saida -----------------------
def _sync_nf(dias):
    fim = _agora()
    ini = (fim - dt.timedelta(days=dias)).replace(hour=0, minute=0, second=0)
    nfs, pagina = [], 1
    while True:
        r = _chamar("/nfe", {"tipo": 1, "pagina": pagina, "limite": 100,
                             "dataEmissaoInicial": ini.strftime("%Y-%m-%d %H:%M:%S"),
                             "dataEmissaoFinal": fim.strftime("%Y-%m-%d 23:59:59")})
        dados = (r or {}).get("data", []) or []
        for n in dados:
            try:
                sit = int(n.get("situacao"))
            except (TypeError, ValueError):
                sit = 0
            if sit in SIT_NF_OK:
                nfs.append({"id": str(n["id"]), "data": str(n.get("dataEmissao", ""))[:10],
                            "natureza": str((n.get("naturezaOperacao") or {}).get("id", ""))})
        if len(dados) < 100:
            break
        pagina += 1
        time.sleep(PAUSA)
    cache = CACHE["nf"]
    pend = [n for n in nfs if n["id"] not in cache]
    with _lock:
        JOB.update(total=len(nfs), feitas=len(nfs) - len(pend))
    for i, n in enumerate(pend, 1):
        time.sleep(PAUSA)
        cache[n["id"]] = {"data": n["data"], "natureza": n["natureza"],
                          "itens": _itens(_chamar("/nfe/%s" % n["id"]))}
        with _lock:
            JOB["feitas"] += 1
        if i % 300 == 0:
            _salvar_cache("nf")
    ids = {n["id"] for n in nfs}
    for k in list(cache):
        if k not in ids:
            del cache[k]
    _salvar_cache("nf")


# --------------------- sincronizacao: pedidos de venda do Full ---------------
def lojas_full():
    txt = kv_get("lojas_full", None)
    txt = FULL_PADRAO if txt is None else txt
    return [x.strip() for x in str(txt).split(",") if x.strip()]


def _sync_full(dias):
    lojas = set(lojas_full())
    if not lojas:
        raise RuntimeError("Informe o ID da loja do Full (veja as lojas encontradas na aba Full).")
    fim = _agora()
    ini = fim - dt.timedelta(days=dias)
    lista, pagina = [], 1
    while True:
        r = _chamar("/pedidos/vendas", {"pagina": pagina, "limite": 100,
                                        "dataInicial": ini.strftime("%Y-%m-%d"),
                                        "dataFinal": fim.strftime("%Y-%m-%d")})
        dados = (r or {}).get("data", []) or []
        for p in dados:
            lista.append({"id": str(p["id"]), "data": str(p.get("data", ""))[:10],
                          "loja": str((p.get("loja") or {}).get("id", "")),
                          "sit": (p.get("situacao") or {}).get("valor")})
        if len(dados) < 100:
            break
        pagina += 1
        time.sleep(PAUSA)
    LOJAS_VISTAS.clear()
    LOJAS_VISTAS.update(Counter(p["loja"] for p in lista))
    alvo = [p for p in lista if p["loja"] in lojas and p["sit"] != SIT_PEDIDO_CANCELADO]
    cache = CACHE["full"]
    pend = [p for p in alvo if p["id"] not in cache]
    with _lock:
        JOB.update(total=len(alvo), feitas=len(alvo) - len(pend))
    for i, p in enumerate(pend, 1):
        time.sleep(PAUSA)
        cache[p["id"]] = {"data": p["data"], "loja": p["loja"],
                          "itens": _itens(_chamar("/pedidos/vendas/%s" % p["id"]))}
        with _lock:
            JOB["feitas"] += 1
        if i % 300 == 0:
            _salvar_cache("full")
    ids = {p["id"] for p in alvo}
    for k in list(cache):
        if k not in ids:
            del cache[k]
    _salvar_cache("full")


def _rodar(fonte, dias):
    try:
        (_sync_nf if fonte == "nf" else _sync_full)(dias)
        with _lock:
            JOB.update(estado="pronto", erro="")
    except Exception as e:
        with _lock:
            JOB.update(estado="erro", erro=str(e)[:300])


@app.route("/api/sincronizar", methods=["POST"])
@exige_login
def api_sincronizar():
    b = request.get_json(silent=True) or {}
    fonte = b.get("fonte")
    if fonte not in ("nf", "full"):
        return jsonify(ok=False, erro="Fonte inválida"), 400
    if fonte == "full" and "lojas" in b:
        kv_set("lojas_full", ",".join(re.findall(r"\d+", str(b["lojas"]))))
    dias = max(7, min(int(b.get("dias", 30)), 120))
    with _lock:
        if JOB["estado"] == "rodando":
            return jsonify(ok=False, erro="Já existe uma atualização em andamento."), 409
        JOB.update(fonte=fonte, estado="rodando", feitas=0, total=0, erro="")
    threading.Thread(target=_rodar, args=(fonte, dias), daemon=True).start()
    return jsonify(ok=True)


@app.route("/api/status")
@exige_login
def api_status():
    nat = Counter(v["natureza"] for v in CACHE["nf"].values())
    with _lock:
        job = dict(JOB)
    return jsonify(
        conectado=bool(_TOK.get("refresh_token") or (kv_get("bling_token") or {}).get("refresh_token")),
        job=job,
        nf={"nfs": len(CACHE["nf"]), "naturezas": [{"id": k, "nfs": c} for k, c in nat.most_common()]},
        full={"pedidos": len(CACHE["full"]), "lojas": ",".join(lojas_full()),
              "lojas_vistas": [{"id": k, "pedidos": c} for k, c in Counter(LOJAS_VISTAS).most_common()]},
        planilha={"itens": len(PLAN), "nome": PLAN_NOME})


# =============================================================================
# Planilha do fornecedor
# =============================================================================
_RE_COD = re.compile(r"^\s*([A-Za-z]{1,3}\d{3,})(?:\s*\(\s*(\d+)\s*\)?)?")


def _chave(m):
    return m.group(1).upper() + ("(%s)" % m.group(2) if m.group(2) else "")


def _num(v):
    if isinstance(v, (int, float)):
        return int(v)
    s = re.sub(r"[^\d,.\-]", "", str(v or ""))
    if not s or s in ("-", ".", ","):
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return int(float(s))
    except ValueError:
        return None


def _linhas_arquivo(nome, conteudo):
    nome = (nome or "").lower()
    if nome.endswith(".xls"):
        try:
            import xlrd
        except ImportError:
            raise ValueError("Falta a biblioteca xlrd no requirements.txt")
        sh = xlrd.open_workbook(file_contents=conteudo).sheet_by_index(0)
        return [sh.row_values(i) for i in range(sh.nrows)]
    if nome.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    texto = conteudo.decode("utf-8", errors="ignore") if isinstance(conteudo, bytes) else conteudo
    return [re.split(r"\t|;|\s{2,}", l.strip()) for l in texto.splitlines() if l.strip()]


def ler_planilha(nome, conteudo):
    out = {}
    for r in _linhas_arquivo(nome, conteudo):
        cels = ["" if c is None else (str(int(c)) if isinstance(c, float) and c == int(c) else str(c)) for c in r]
        m = _RE_COD.match(cels[0]) if cels else None
        if not m or len(cels) < 2:
            continue
        saldo = _num(cels[-1])
        if saldo is None:
            continue
        out[_chave(m)] = {"desc": cels[1].lstrip(") ").strip() if len(cels) > 2 else "",
                          "apl": cels[2].strip() if len(cels) > 3 else "", "saldo": saldo}
    return out


def ler_texto(texto):
    out = {}
    for l in (texto or "").splitlines():
        m, n = _RE_COD.match(l), re.search(r"(-?\d[\d.]*)\s*$", l)
        if m and n:
            p = re.split(r"\s{2,}|\t", l.strip())
            out[_chave(m)] = {"desc": p[1] if len(p) > 2 else "", "apl": p[2] if len(p) > 3 else "",
                              "saldo": _num(n.group(1)) or 0}
    return out


@app.route("/api/planilha", methods=["POST"])
@exige_login
def api_planilha():
    global PLAN, PLAN_NOME
    try:
        arq = request.files.get("arquivo")
        if arq and arq.filename:
            plan, nome = ler_planilha(arq.filename, arq.read()), arq.filename
        else:
            plan, nome = ler_texto(request.form.get("texto", "")), "texto colado"
    except Exception as e:
        return jsonify(ok=False, erro="Não consegui ler a planilha: %s" % e), 400
    if not plan:
        return jsonify(ok=False, erro="Nenhum produto encontrado. Use o modelo CODIGO / DESCRICAO / APLICACAO / SALDO."), 400
    PLAN, PLAN_NOME = plan, nome
    try:
        kv_set("planilha", plan)
        kv_set("planilha_nome", nome)
    except Exception:
        pass
    return jsonify(ok=True, itens=len(plan), nome=nome)


# =============================================================================
# Analise
# =============================================================================
def _componentes(codigo, conhecidos):
    """Kit 'SM804014-SM804022-PM102' vira cada peca que existe na planilha."""
    if codigo in conhecidos:
        return [codigo]
    return [t for t in re.split(r"[-+/\s]+", codigo) if t in conhecidos]


def analisar(aba, naturezas, modo, mult):
    usa_nf, usa_full = aba in ("nf", "total"), aba in ("full", "total")
    hoje = _agora().date()
    lim = [hoje - dt.timedelta(days=d) for d in (7, 30, 60)]
    q = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0, 0.0])      # 7d, 30d, 60d, NF 30d, Full 30d
    nao_map, docs = Counter(), 0
    fontes = []
    if usa_nf:
        fontes.append(("nf", [v for v in CACHE["nf"].values() if not naturezas or v["natureza"] in naturezas]))
    if usa_full:
        lf = set(lojas_full())
        fontes.append(("full", [v for v in CACHE["full"].values() if not lf or v["loja"] in lf]))
    for nome, docs_f in fontes:
        for d in docs_f:
            try:
                dia = dt.date.fromisoformat(d["data"])
            except ValueError:
                continue
            docs += 1
            for cod, qtd in d["itens"]:
                comps = _componentes(cod, PLAN)
                if not comps:
                    nao_map[cod] += qtd
                    continue
                for c in comps:
                    for i in range(3):
                        if dia > lim[i]:
                            q[c][i] += qtd
                    if dia > lim[1]:
                        q[c][3 if nome == "nf" else 4] += qtd
    linhas = []
    for cod, p in PLAN.items():
        q7, q30, q60, nf30, fu30 = q[cod] if cod in q else (0, 0, 0, 0, 0)
        taxa = (q30 / 30.0 if modo == "30d" else q60 / 60.0 if modo == "60d" else max(q7 / 7.0, q30 / 30.0)) * mult
        saldo = p["saldo"]
        if taxa <= 0:
            risco, cob = "Sem saída", None
        elif saldo <= 0:
            risco, cob = "Zerado", 0.0
        else:
            cob = saldo / taxa
            risco = "Falta em 1 semana" if cob < 7 else ("Falta em 1 mês" if cob < 30 else "OK")
        linhas.append({"codigo": cod, "desc": p["desc"], "apl": p["apl"], "saldo": saldo,
                       "q7": round(q7), "q30": round(q30), "q60": round(q60), "nf30": round(nf30), "full30": round(fu30),
                       "taxa": round(taxa, 2), "cobertura": None if cob is None else round(cob, 1), "risco": risco,
                       "falta30": max(0, math.ceil(taxa * 30 - saldo)) if taxa > 0 else 0})
    ordem = {"Zerado": 0, "Falta em 1 semana": 1, "Falta em 1 mês": 2, "OK": 3, "Sem saída": 4}
    linhas.sort(key=lambda l: (ordem[l["risco"]], l["cobertura"] if l["cobertura"] is not None else 1e9, -l["q30"]))
    resumo = dict(Counter(l["risco"] for l in linhas))
    resumo.update(docs=docs, produtos=len(PLAN))
    return linhas, resumo, [{"codigo": k, "qtd": round(v)} for k, v in nao_map.most_common(15)]


@app.route("/api/analisar", methods=["POST"])
@exige_login
def api_analisar():
    b = request.get_json(silent=True) or {}
    aba = b.get("aba", "nf")
    if not PLAN:
        return jsonify(ok=False, erro="Envie a planilha do fornecedor primeiro."), 400
    if aba in ("nf", "total") and not CACHE["nf"] and aba == "nf":
        return jsonify(ok=False, erro="Atualize as NFs do Bling antes de analisar."), 400
    if aba == "full" and not CACHE["full"]:
        return jsonify(ok=False, erro="Atualize os pedidos do Full antes de analisar."), 400
    if aba == "total" and not (CACHE["nf"] or CACHE["full"]):
        return jsonify(ok=False, erro="Atualize as NFs e/ou os pedidos do Full antes de analisar."), 400
    try:
        mult = max(1.0, float(str(b.get("mult", 1)).replace(",", ".")))
    except ValueError:
        mult = 1.0
    linhas, resumo, nm = analisar(aba, set(b.get("naturezas") or []), b.get("modo", "conservador"), mult)
    ULTIMA.update(linhas=linhas, aba=aba)
    return jsonify(ok=True, linhas=linhas, resumo=resumo, nao_mapeados=nm)


@app.route("/api/exportar.xlsx")
@exige_login
def api_exportar():
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risco de falta"
    ws.append(["Código", "Descrição", "Aplicação", "Saldo fornecedor", "Saída 7d", "Saída 30d", "Saída 60d",
               "NF 30d", "Full 30d", "Saída/dia", "Cobertura (dias)", "Risco", "Faltaria em 30d (un)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    cores = {"Zerado": "F4B6B6", "Falta em 1 semana": "F9D2A8", "Falta em 1 mês": "FBEFA6"}
    for l in ULTIMA["linhas"]:
        ws.append([l["codigo"], l["desc"], l["apl"], l["saldo"], l["q7"], l["q30"], l["q60"], l["nf30"], l["full30"],
                   l["taxa"], l["cobertura"], l["risco"], l["falta30"]])
        if l["risco"] in cores:
            ws.cell(row=ws.max_row, column=12).fill = PatternFill("solid", fgColor=cores[l["risco"]])
    for col, w in zip("ABCDEFGHIJKLM", (12, 38, 40, 14, 10, 10, 10, 10, 10, 10, 14, 20, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="risco_falta_%s_%s.xlsx" % (ULTIMA["aba"], _agora().strftime("%Y%m%d")),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# =============================================================================
# Pagina
# =============================================================================
PAGINA = r"""<!doctype html>
<html lang="pt-br"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Risco de falta</title>
<style>
 body{font:15px/1.45 system-ui,sans-serif;margin:0;background:#f5f6f8;color:#1c2330}
 main{max-width:1200px;margin:0 auto;padding:20px}
 header{display:flex;justify-content:space-between;align-items:baseline}
 h1{font-size:22px;margin:0 0 4px} p.sub{margin:0 0 14px;color:#5b6577} a{color:#1f5eff}
 section{background:#fff;border:1px solid #dfe3ea;border-radius:8px;padding:16px;margin-bottom:14px}
 h2{font-size:16px;margin:0 0 10px}
 .abas{display:flex;gap:6px;margin-bottom:-1px}
 .abas button{background:#e8ecf4;color:#1c2330;border-radius:8px 8px 0 0;padding:9px 16px}
 .abas button.on{background:#fff;font-weight:600;border:1px solid #dfe3ea;border-bottom-color:#fff}
 .painel{border-top-left-radius:0}
 button{background:#1f5eff;color:#fff;border:0;border-radius:6px;padding:9px 14px;font-size:14px;cursor:pointer}
 button.sec{background:#e8ecf4;color:#1c2330} button:disabled{opacity:.5;cursor:default}
 label{margin-right:14px;display:inline-block;margin-bottom:6px} select,input[type=number],input[type=text]{padding:6px;border:1px solid #c9cfda;border-radius:5px}
 textarea{width:100%;height:60px;border:1px solid #c9cfda;border-radius:5px;padding:6px;font:12px monospace;box-sizing:border-box}
 progress{width:260px;vertical-align:middle} .muted{color:#5b6577;font-size:13px}
 .cards{display:flex;gap:10px;flex-wrap:wrap;margin:10px 0}
 .card{flex:1;min-width:150px;border-radius:8px;padding:10px 12px;cursor:pointer;border:2px solid transparent}
 .card b{font-size:24px;display:block} .card.on{border-color:#1c2330}
 .c0{background:#f4b6b6}.c1{background:#f9d2a8}.c2{background:#fbefa6}.c3{background:#cfe9d4}.c4{background:#e3e6ec}
 table{border-collapse:collapse;width:100%;font-size:13px} th,td{padding:6px 8px;border-bottom:1px solid #edf0f5;text-align:left}
 th{cursor:pointer;background:#f0f2f7;position:sticky;top:0} td.n,th.n{text-align:right}
 .wrap{max-height:560px;overflow:auto;border:1px solid #dfe3ea;border-radius:6px}
 .erro{color:#b3261e;font-weight:600} .oculto{display:none}
</style></head><body><main>
<header><div><h1>Risco de falta</h1>
<p class="sub">Saídas do Bling cruzadas com o saldo da planilha do fornecedor.</p></div><a href="/logout">Sair</a></header>

<section id="conexao" class="oculto"><b>O Bling ainda não está conectado.</b> <a href="/bling/conectar"><button>Conectar ao Bling</button></a></section>

<section>
 <h2>Planilha de estoque do fornecedor</h2>
 <input type="file" id="arq" accept=".xls,.xlsx,.xlsm,.csv,.txt"> <button class="sec" id="bPlan">Carregar planilha</button>
 <span id="planTxt" class="muted"></span>
 <p class="muted" style="margin:8px 0 4px">Ou cole as linhas (CODIGO, DESCRICAO, APLICACAO, SALDO):</p>
 <textarea id="texto"></textarea>
 <div id="planMsg" class="erro"></div>
</section>

<div class="abas">
 <button data-aba="nf" class="on">Saídas por NF</button>
 <button data-aba="full">Full (pedidos de venda)</button>
 <button data-aba="total">Total (NF + Full)</button>
</div>
<section class="painel">
 <div id="p-nf">
  <label>Período <select id="diasNf"><option>30</option><option selected>60</option><option>90</option></select> dias</label>
  <button id="bNf">Atualizar NFs</button>
  <div class="muted" id="nfTxt" style="margin:8px 0"></div>
  <div id="nat"></div>
 </div>
 <div id="p-full" class="oculto">
  <label>ID(s) da loja do Full <input type="text" id="lojas" size="22"></label>
  <label>Período <select id="diasFull"><option selected>30</option><option>60</option></select> dias</label>
  <button id="bFull">Atualizar pedidos do Full</button>
  <div class="muted" id="fullTxt" style="margin:8px 0"></div>
  <div class="muted" id="vistas"></div>
 </div>
 <div id="p-total" class="oculto muted">Soma as NFs de saída com os pedidos do Full já atualizados nas outras duas abas. Atualize cada uma antes, se precisar de dados novos.</div>
 <div id="progWrap" class="oculto" style="margin:8px 0"><progress id="prog" value="0" max="1"></progress> <span id="progTxt" class="muted"></span></div>
 <div id="jobErro" class="erro"></div>
 <hr style="border:0;border-top:1px solid #edf0f5;margin:14px 0">
 <label>Ritmo de saída
  <select id="modo"><option value="conservador">Conservador (maior entre 7 e 30 dias)</option><option value="30d">Média dos últimos 30 dias</option><option value="60d">Média dos últimos 60 dias</option></select></label>
 <label>Multiplicador de demanda <input type="number" id="mult" value="1" step="0.1" min="1" style="width:70px"></label>
 <button id="bAn">Analisar</button> <button class="sec oculto" id="bXls">Baixar Excel</button>
 <div class="muted">O multiplicador cobre quem compra do mesmo fornecedor e não aparece neste Bling (2 = dobra a demanda).</div>
 <div id="msg" class="erro"></div>
</section>

<section id="res" class="oculto">
 <h2 id="resTitulo">Resultado</h2>
 <div class="cards" id="cards"></div>
 <div class="muted" id="info"></div>
 <div class="wrap"><table><thead><tr id="cab"></tr></thead><tbody id="corpo"></tbody></table></div>
 <div id="nm" class="muted" style="margin-top:10px"></div>
</section>

<script>
const $=id=>document.getElementById(id);
const COLS=[["codigo","Código"],["desc","Descrição"],["apl","Aplicação"],["saldo","Saldo forn.",1],["q7","7d",1],["q30","30d",1],["q60","60d",1],["nf30","NF 30d",1,"total"],["full30","Full 30d",1,"total"],["taxa","Saída/dia",1],["cobertura","Cobertura (dias)",1],["risco","Risco"],["falta30","Faltaria em 30d",1]];
const RISCOS=["Zerado","Falta em 1 semana","Falta em 1 mês","OK","Sem saída"];
const TITULOS={nf:"Saídas por NF",full:"Full (pedidos de venda)",total:"Total (NF + Full)"};
let ABA="nf",DADOS=[],filtro="",ordem=null,asc=true,poll=null,jaPreencheu=false,ABA_RES="nf";

document.querySelectorAll(".abas button").forEach(b=>b.onclick=()=>{
 ABA=b.dataset.aba;document.querySelectorAll(".abas button").forEach(x=>x.classList.toggle("on",x===b));
 ["nf","full","total"].forEach(a=>$("p-"+a).classList.toggle("oculto",a!==ABA));
});
async function api(url,opt){const r=await fetch(url,opt);if(r.status===401){location.href="/login";throw 0;}return r.json();}

async function status(){
 const r=await api("/api/status"),j=r.job;
 $("conexao").classList.toggle("oculto",r.conectado);
 $("planTxt").textContent=r.planilha.itens?r.planilha.itens+" produtos carregados ("+r.planilha.nome+")":"";
 $("nfTxt").textContent=r.nf.nfs+" NFs de saída em memória";
 $("fullTxt").textContent=r.full.pedidos+" pedidos do Full em memória";
 if(!jaPreencheu){$("lojas").value=r.full.lojas;jaPreencheu=true;}
 $("vistas").textContent=r.full.lojas_vistas.length?"Lojas encontradas no Bling na última atualização (pedidos): "+r.full.lojas_vistas.map(x=>x.id+" ("+x.pedidos+")").join(" · "):"";
 const marc=[...document.querySelectorAll(".nt:checked")].map(x=>x.value);
 if(r.nf.naturezas.length)$("nat").innerHTML='<div class="muted">Natureza de operação a considerar (desmarque remessas ao Full e transferências):</div>'+r.nf.naturezas.map((n,i)=>
  '<label><input type="checkbox" class="nt" value="'+n.id+'" '+((marc.length?marc.includes(n.id):i===0)?"checked":"")+'> '+n.id+" ("+n.nfs+" NFs)</label>").join("");
 const rod=j.estado==="rodando";
 $("progWrap").classList.toggle("oculto",!rod);$("bNf").disabled=$("bFull").disabled=rod;
 if(rod){$("prog").max=Math.max(j.total,1);$("prog").value=j.feitas;$("progTxt").textContent=(j.fonte==="nf"?"NFs: ":"Pedidos: ")+j.feitas+" / "+j.total;}
 $("jobErro").textContent=j.estado==="erro"?"Erro: "+j.erro:"";
 if(!rod&&poll){clearInterval(poll);poll=null;}
}
async function sync(fonte){
 const body={fonte,dias:+(fonte==="nf"?$("diasNf"):$("diasFull")).value};
 if(fonte==="full")body.lojas=$("lojas").value;
 const j=await api("/api/sincronizar",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(body)});
 if(!j.ok){$("jobErro").textContent=j.erro;return;}
 if(!poll)poll=setInterval(status,1500);status();
}
$("bNf").onclick=()=>sync("nf");$("bFull").onclick=()=>sync("full");
$("bPlan").onclick=async()=>{
 $("planMsg").textContent="";const fd=new FormData();
 if($("arq").files[0])fd.append("arquivo",$("arq").files[0]);fd.append("texto",$("texto").value);
 const j=await api("/api/planilha",{method:"POST",body:fd});
 if(!j.ok){$("planMsg").textContent=j.erro;return;}status();
};
$("bAn").onclick=async()=>{
 $("msg").textContent="";$("bAn").disabled=true;
 try{
  const j=await api("/api/analisar",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify({
   aba:ABA,modo:$("modo").value,mult:$("mult").value,naturezas:[...document.querySelectorAll(".nt:checked")].map(x=>x.value)})});
  if(!j.ok){$("msg").textContent=j.erro;return;}
  DADOS=j.linhas;filtro="";ABA_RES=ABA;mostrar(j);
 }finally{$("bAn").disabled=false;}
};
$("bXls").onclick=()=>{location.href="/api/exportar.xlsx"};
function mostrar(j){
 $("res").classList.remove("oculto");$("bXls").classList.remove("oculto");$("resTitulo").textContent="Resultado: "+TITULOS[ABA_RES];
 $("cards").innerHTML=RISCOS.map((r,i)=>'<div class="card c'+i+'" data-r="'+r+'"><b>'+(j.resumo[r]||0)+'</b>'+r+"</div>").join("");
 document.querySelectorAll(".card").forEach(c=>c.onclick=()=>{filtro=filtro===c.dataset.r?"":c.dataset.r;desenhar();});
 $("info").textContent=j.resumo.produtos+" produtos na planilha · "+j.resumo.docs+" documentos considerados";
 $("nm").textContent=j.nao_mapeados.length?"Vendidos fora da planilha (outras marcas, por exemplo): "+j.nao_mapeados.map(x=>x.codigo+" ("+x.qtd+")").join(", "):"";
 const cols=COLS.filter(c=>!c[3]||c[3]===ABA_RES);
 $("cab").innerHTML=cols.map(c=>'<th class="'+(c[2]?"n":"")+'" data-k="'+c[0]+'">'+c[1]+"</th>").join("");
 document.querySelectorAll("#cab th").forEach(th=>th.onclick=()=>{const k=th.dataset.k;asc=ordem===k?!asc:true;ordem=k;desenhar();});
 desenhar();
}
function desenhar(){
 document.querySelectorAll(".card").forEach(c=>c.classList.toggle("on",c.dataset.r===filtro));
 const cols=COLS.filter(c=>!c[3]||c[3]===ABA_RES);
 let l=DADOS.filter(x=>!filtro||x.risco===filtro);
 if(ordem)l=[...l].sort((a,b)=>{let x=a[ordem],y=b[ordem];if(x===null)x=1e12;if(y===null)y=1e12;return (x>y?1:x<y?-1:0)*(asc?1:-1)});
 $("corpo").innerHTML=l.slice(0,1500).map(x=>"<tr>"+cols.map(c=>'<td class="'+(c[2]?"n":"")+'">'+(x[c[0]]===null?"—":String(x[c[0]]).replace(/</g,"&lt;"))+"</td>").join("")+"</tr>").join("");
}
status();
</script></main></body></html>"""


@app.route("/")
@exige_login
def pagina():
    return Response(PAGINA, mimetype="text/html")
